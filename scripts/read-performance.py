"""Read the original workbook without saving or changing any of its contents."""
import colorsys
import json
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel, WINDOWS_EPOCH


def export_workbook(path):
    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    theme = {}
    if book.loaded_theme:
        xml = ElementTree.fromstring(book.loaded_theme)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = xml.find("a:themeElements/a:clrScheme", ns)
        if scheme is not None:
            for child in scheme:
                color = next(iter(child), None)
                if color is not None:
                    theme[child.tag.split("}")[-1]] = color.get("lastClr") or color.get("val")
    theme_keys = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]

    def color_hex(color, default="000000"):
        value = default
        if color is not None:
            if color.type == "rgb":
                value = color.rgb[-6:]
            elif color.type == "indexed" and 0 <= color.indexed < len(COLOR_INDEX):
                value = COLOR_INDEX[color.indexed][-6:]
            elif color.type == "theme" and 0 <= color.theme < len(theme_keys):
                value = theme.get(theme_keys[color.theme], default)
            tint = color.tint or 0
            if tint:
                rgb = [int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)]
                hue, light, saturation = colorsys.rgb_to_hls(*rgb)
                light = light * (1 + tint) if tint < 0 else light * (1 - tint) + tint
                value = "".join(f"{round(c * 255):02x}" for c in colorsys.hls_to_rgb(hue, light, saturation))
        return "#" + value

    def cell_style(cell):
        font, alignment, fill, border = cell.font, cell.alignment, cell.fill, cell.border
        css = {
            "font-family": json.dumps(font.name or "Calibri", ensure_ascii=False) + ", Arial, sans-serif",
            "font-size": f"{font.sz or 11}pt",
            "font-weight": "700" if font.b else "400",
            "font-style": "italic" if font.i else "normal",
            "color": color_hex(font.color),
            "text-align": {"general": "right" if cell.data_type in ("n", "f") and cell.value is not None else "left", "centerContinuous": "center", "fill": "left", "distributed": "justify"}.get(alignment.horizontal or "general", alignment.horizontal),
            "vertical-align": {"center": "middle", "distributed": "middle", "justify": "middle"}.get(alignment.vertical or "bottom", alignment.vertical or "bottom"),
        }
        if fill.patternType == "solid":
            css["background-color"] = color_hex(fill.fgColor, "FFFFFF")
        decorations = []
        if font.u:
            decorations.append("underline")
        if font.strike:
            decorations.append("line-through")
        if decorations:
            css["text-decoration"] = " ".join(decorations)
        if alignment.indent:
            css["padding-left"] = f"{alignment.indent * 3}ch"
        for edge in ["top", "right", "bottom", "left"]:
            side = getattr(border, edge)
            if not side or not side.style:
                continue
            kind = side.style.lower()
            weight = 3 if "thick" in kind else 2 if "medium" in kind else 1
            pattern = "double" if "double" in kind else "dotted" if "dot" in kind else "dashed" if "dash" in kind else "solid"
            css[f"border-{edge}"] = f"{weight}px {pattern} {color_hex(side.color)}"
        return css

    sheets = []
    for sheet in book.worksheets:
        spans, covered = {}, set()
        for merge in sheet.merged_cells.ranges:
            spans[(merge.min_row, merge.min_col)] = [merge.max_row - merge.min_row + 1, merge.max_col - merge.min_col + 1]
            covered.update((r, c) for r in range(merge.min_row, merge.max_row + 1) for c in range(merge.min_col, merge.max_col + 1) if (r, c) != (merge.min_row, merge.min_col))
        columns = []
        for index in range(1, sheet.max_column + 1):
            width = sheet.sheet_format.defaultColWidth or 8.43
            for dimension in sheet.column_dimensions.values():
                if dimension.min <= index <= dimension.max:
                    width = dimension.width
                    break
            columns.append(round(width * 7 + 5, 2))
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            cells = []
            for cell in row:
                if (cell.row, cell.column) in covered:
                    continue
                value = cell.value
                is_formula = cell.data_type == "f"
                if is_formula:
                    value = cached[sheet.title][cell.coordinate].value
                    if value is None:
                        raise ValueError(f"{sheet.title}!{cell.coordinate} has no saved formula result. Recalculate and save in Excel before replacing the workbook.")
                if isinstance(value, (datetime, date, time, timedelta)):
                    if isinstance(value, timedelta):
                        value = value.total_seconds() / 86400
                    else:
                        value = to_excel(value, book.epoch)
                kind = "blank" if value is None else "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float)) else "text"
                cells.append({"address": cell.coordinate, "value": value, "kind": kind, "format": cell.number_format, "style": cell_style(cell), "span": spans.get((cell.row, cell.column), [1, 1]), "link": cell.hyperlink.target if cell.hyperlink else None, "formula": is_formula})
            row_index = row[0].row
            rows.append({"height": sheet.row_dimensions[row_index].height or sheet.sheet_format.defaultRowHeight or 15, "cells": cells})
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "widths": columns, "data": rows})
    return {"date1904": book.epoch != WINDOWS_EPOCH, "sheets": sheets}


if __name__ == "__main__":
    print(json.dumps(export_workbook(Path(sys.argv[1])), ensure_ascii=False))
