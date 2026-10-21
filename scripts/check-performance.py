"""Independently check generated HTML against every original workbook cell."""
import json
import sys
from datetime import date, datetime, time, timedelta
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel


class Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.sheets = {}
        self.sheet = self.cell = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "article" and "data-sheet-name" in attrs:
            self.sheet = {"attrs": attrs, "cells": {}, "rows": 0, "columns": 0}
            self.sheets[attrs["data-sheet-name"]] = self.sheet
        elif self.sheet is not None:
            if tag == "tr":
                self.sheet["rows"] += 1
            elif tag == "col":
                self.sheet["columns"] += 1
            elif tag == "td":
                self.cell = {"attrs": attrs, "text": ""}
                self.sheet["cells"][attrs["data-cell"]] = self.cell

    def handle_data(self, text):
        if self.cell is not None:
            self.cell["text"] += text

    def handle_endtag(self, tag):
        if tag == "td":
            self.cell = None
        elif tag == "article":
            self.sheet = None


def check(source, index):
    book = load_workbook(source, data_only=False)
    results = load_workbook(source, data_only=True)
    html = Tables()
    html.feed(Path(index).read_text())
    assert list(html.sheets) == book.sheetnames, "Sheet names or order changed"
    count = 0
    for sheet in book:
        actual = html.sheets[sheet.title]
        assert actual["rows"] == sheet.max_row, f"{sheet.title}: row count changed"
        assert actual["columns"] == sheet.max_column, f"{sheet.title}: column count changed"
        merged = {}
        covered = set()
        for span in sheet.merged_cells.ranges:
            merged[sheet.cell(span.min_row, span.min_col).coordinate] = (span.max_row - span.min_row + 1, span.max_col - span.min_col + 1)
            covered.update(sheet.cell(r, c).coordinate for r in range(span.min_row, span.max_row + 1) for c in range(span.min_col, span.max_col + 1) if (r, c) != (span.min_row, span.min_col))
        for row in sheet:
            for cell in row:
                if cell.coordinate in covered:
                    continue
                exported = actual["cells"][cell.coordinate]
                value = results[sheet.title][cell.coordinate].value if cell.data_type == "f" else cell.value
                if isinstance(value, timedelta):
                    value = value.total_seconds() / 86400
                elif isinstance(value, (date, datetime, time)):
                    value = to_excel(value, book.epoch)
                assert json.loads(exported["attrs"]["data-raw-value"]) == value, f"{sheet.title}!{cell.coordinate}: value changed"
                assert exported["attrs"]["data-number-format"] == cell.number_format, f"{cell.coordinate}: format changed"
                if isinstance(value, str):
                    assert exported["text"] == value, f"{cell.coordinate}: text changed"
                expected_span = merged.get(cell.coordinate, (1, 1))
                assert (int(exported["attrs"].get("rowspan", 1)), int(exported["attrs"].get("colspan", 1))) == expected_span, f"{cell.coordinate}: merge changed"
                count += 1
    print(f"Verified {len(book.sheetnames)} sheets and {count} cells against the original workbook.")


if __name__ == "__main__":
    check(sys.argv[1], sys.argv[2])
